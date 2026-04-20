from .local_rag import LocalRAG
from .vector_rag import VectorRAG
import os
import re
import json
import tempfile
from datetime import datetime
from typing import List, Dict, Optional, Tuple
from urllib.parse import urlparse
from enum import Enum
from utils.logger import logger


class ErrorCode(Enum):
    WEB_PARSER_UNAVAILABLE = "E001"
    SELENIUM_UNAVAILABLE = "E002"
    API_TIMEOUT = "E003"
    API_UNAVAILABLE = "E004"
    PDF_PARSE_ERROR = "E005"
    DOCX_PARSE_ERROR = "E006"
    XLSX_PARSE_ERROR = "E007"
    HTML_PARSE_ERROR = "E008"
    INVALID_URL = "E009"
    CONFIG_MISSING = "E010"


# 依赖检查
try:
    import requests
    from bs4 import BeautifulSoup
    HAS_WEB_PARSER = True
except ImportError:
    HAS_WEB_PARSER = False
    logger.warning("requests 或 BeautifulSoup 未安装，网页解析功能不可用")

try:
    import PyPDF2
    HAS_PDF = True
except ImportError:
    HAS_PDF = False

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    import openpyxl
    HAS_XLSX = True
except ImportError:
    HAS_XLSX = False

try:
    from selenium import webdriver
    from selenium.webdriver.chrome.options import Options as ChromeOptions
    from selenium.webdriver.edge.options import Options as EdgeOptions
    from selenium.webdriver.common.by import By
    from selenium.webdriver.support.ui import WebDriverWait
    from selenium.webdriver.support import expected_conditions as EC
    HAS_SELENIUM = True
except ImportError:
    HAS_SELENIUM = False
    logger.warning("Selenium 未安装，动态网页解析功能不可用。请运行: pip install selenium webdriver-manager")

try:
    from openai import OpenAI
    from dotenv import load_dotenv
    load_dotenv()
    API_KEY = os.getenv("API_KEY")
    MODEL_ID = os.getenv("MODEL_ID")
    
    if not API_KEY or not MODEL_ID:
        logger.error("API_KEY 或 MODEL_ID 未配置，请在 .env 文件中设置")
        HAS_API = False
    else:
        HAS_API = True
        logger.info("API 依赖已加载")
except ImportError as e:
    logger.warning(f"API 依赖未安装，将使用备用模式运行: {e}")
    HAS_API = False

# 初始化 OpenAI 客户端
if HAS_API:
    client = OpenAI(
        base_url="https://ark.cn-beijing.volces.com/api/v1",
        api_key=API_KEY,
    )


class WebParser:
    """网页解析器"""
    
    SUPPORTED_FORMATS = ['html', 'pdf', 'txt', 'csv', 'docx', 'doc', 'xlsx', 'xls', 'pptx', 'ppt', 'md', 'mobi', 'epub']
    MAX_URLS_PER_CALL = 3
    DYNAMIC_PAGE_INDICATORS = [
        '需要允许', '请启用', 'JavaScript', 'javascript', 'loading', 
        '请开启', '浏览器不支持', 'enable javascript', 'enable-js'
    ]
    
    def __init__(self):
        self.session = self._init_session()
        self._driver = None
    
    def _init_session(self):
        """初始化请求会话"""
        if not HAS_WEB_PARSER:
            return None
        session = requests.Session()
        session.headers.update({
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
        })
        return session
    
    def _get_driver(self):
        """获取浏览器驱动"""
        if not HAS_SELENIUM:
            return None
        if self._driver is None:
            driver = self._try_edge_driver() or self._try_chrome_driver()
            if driver:
                self._driver = driver
                return driver
            logger.warning("未找到可用的浏览器 (Edge/Chrome)")
            return None
        return self._driver
    
    def _try_edge_driver(self):
        """尝试使用Edge浏览器"""
        try:
            options = EdgeOptions()
            self._configure_driver_options(options)
            options.add_argument('--user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36 Edg/120.0.0.0')
            driver = webdriver.Edge(options=options)
            driver.set_page_load_timeout(60)
            logger.info("使用 Edge 浏览器解析")
            return driver
        except Exception as e:
            logger.debug(f"Edge 浏览器不可用: {str(e)[:80]}")
            return None
    
    def _try_chrome_driver(self):
        """尝试使用Chrome浏览器"""
        try:
            options = ChromeOptions()
            self._configure_driver_options(options)
            options.add_argument('--user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36')
            driver = webdriver.Chrome(options=options)
            driver.set_page_load_timeout(60)
            logger.info("使用 Chrome 浏览器解析")
            return driver
        except Exception as e:
            logger.debug(f"Chrome 浏览器不可用: {str(e)[:80]}")
            return None
    
    def _configure_driver_options(self, options):
        """配置浏览器选项"""
        options.add_argument('--headless=new')
        options.add_argument('--disable-gpu')
        options.add_argument('--no-sandbox')
        options.add_argument('--disable-dev-shm-usage')
        options.add_argument('--disable-blink-features=AutomationControlled')
        options.add_experimental_option('excludeSwitches', ['enable-automation'])
        options.add_experimental_option('useAutomationExtension', False)
        options.add_argument('--disable-extensions')
        options.add_argument('--disable-infobars')
    
    def _is_dynamic_page(self, content: str) -> bool:
        """判断是否为动态页面"""
        return any(indicator in content for indicator in self.DYNAMIC_PAGE_INDICATORS)
    
    def _parse_with_selenium(self, url: str) -> Dict:
        """使用Selenium解析动态页面"""
        driver = self._get_driver()
        if driver is None:
            return {"url": url, "error": "Selenium 不可用，无法解析动态网页", "error_code": ErrorCode.SELENIUM_UNAVAILABLE.value}
        
        try:
            driver.get(url)
            WebDriverWait(driver, 10).until(
                EC.presence_of_element_located((By.TAG_NAME, "body"))
            )
            import time
            time.sleep(3)
            
            title = driver.title if driver.title else "无标题"
            body = driver.find_element(By.TAG_NAME, "body")
            text = body.text
            
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            content = '\n'.join(lines[:300])
            
            return {
                "url": url,
                "title": title,
                "content": content,
                "format": "html_dynamic"
            }
            
        except Exception as e:
            return {"url": url, "error": f"Selenium解析失败: {str(e)}"}
    
    def close_driver(self):
        """关闭浏览器驱动"""
        if self._driver:
            try:
                self._driver.quit()
            except:
                pass
            self._driver = None
    
    def parse_urls(self, url_list: List[str]) -> List[Dict]:
        """解析URL列表"""
        if not HAS_WEB_PARSER:
            return [{"error": "网页解析功能不可用，请安装 requests 和 beautifulsoup4", "error_code": ErrorCode.WEB_PARSER_UNAVAILABLE.value}]
        
        url_list = url_list[:self.MAX_URLS_PER_CALL]
        return [self.parse_single_url(url) for url in url_list]
    
    def parse_single_url(self, url: str) -> Dict:
        """解析单个URL"""
        if not HAS_WEB_PARSER:
            return {"url": url, "error": "网页解析功能不可用", "error_code": ErrorCode.WEB_PARSER_UNAVAILABLE.value}
        
        try:
            parsed_url = urlparse(url)
            file_ext = parsed_url.path.split('.')[-1].lower() if '.' in parsed_url.path else 'html'
            
            parsers = {
                'pdf': self._parse_pdf_url,
                'txt': lambda u: self._parse_text_url(u, 'txt'),
                'md': lambda u: self._parse_text_url(u, 'md'),
                'csv': lambda u: self._parse_text_url(u, 'csv'),
                'docx': self._parse_docx_url,
                'doc': self._parse_docx_url,
                'xlsx': self._parse_xlsx_url,
                'xls': self._parse_xlsx_url,
            }
            
            parser = parsers.get(file_ext, self._parse_html_url)
            return parser(url)
            
        except Exception as e:
            return {"url": url, "error": str(e)}
    
    def _parse_html_url(self, url: str) -> Dict:
        """解析HTML URL"""
        try:
            response = self.session.get(url, timeout=30)
            response.raise_for_status()
            
            soup = BeautifulSoup(response.content, 'html.parser')
            
            for tag in ["script", "style", "nav", "footer", "header"]:
                for element in soup.find_all(tag):
                    element.decompose()
            
            title = soup.title.string.strip() if soup.title and soup.title.string else "无标题"
            
            meta_desc = soup.find('meta', attrs={'name': 'description'})
            description = meta_desc['content'] if meta_desc and meta_desc.get('content') else ""
            
            main_content = soup.find('main') or soup.find('article') or soup.find('div', class_='content') or soup.find('body')
            text = main_content.get_text(separator='\n', strip=True) if main_content else soup.get_text(separator='\n', strip=True)
            
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            content = '\n'.join(lines[:200])
            
            if self._is_dynamic_page(content) or len(content) < 100:
                if HAS_SELENIUM:
                    logger.info(f"检测到动态网页，使用 Selenium 解析: {url}")
                    return self._parse_with_selenium(url)
            
            return {
                "url": url,
                "title": title,
                "description": description,
                "content": content,
                "format": "html"
            }
            
        except Exception as e:
            if HAS_SELENIUM:
                logger.info(f"静态解析失败，尝试使用 Selenium: {url}")
                return self._parse_with_selenium(url)
            return {"url": url, "error": f"HTML解析失败: {str(e)}"}
    
    def _parse_pdf_url(self, url: str) -> Dict:
        """解析PDF URL"""
        if not HAS_PDF:
            return {"url": url, "error": "PDF解析功能不可用，请安装 PyPDF2", "error_code": ErrorCode.PDF_PARSE_ERROR.value}
        
        try:
            response = self.session.get(url, timeout=60)
            response.raise_for_status()
            
            with tempfile.NamedTemporaryFile(delete=False, suffix='.pdf') as tmp_file:
                tmp_file.write(response.content)
                tmp_path = tmp_file.name
            
            content = ""
            with open(tmp_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages[:20]:
                    content += page.extract_text() + "\n"
            
            os.unlink(tmp_path)
            
            return {
                "url": url,
                "title": url.split('/')[-1],
                "content": content[:10000],
                "format": "pdf"
            }
            
        except Exception as e:
            return {"url": url, "error": f"PDF解析失败: {str(e)}"}
    
    def _parse_text_url(self, url: str, file_ext: str) -> Dict:
        """解析文本URL"""
        try:
            response = self.session.get(url, timeout=30)
            response.raise_for_status()
            
            content = response.text[:10000]
            title = url.split('/')[-1]
            
            return {
                "url": url,
                "title": title,
                "content": content,
                "format": file_ext
            }
            
        except Exception as e:
            return {"url": url, "error": f"文本解析失败: {str(e)}"}
    
    def _parse_docx_url(self, url: str) -> Dict:
        """解析DOCX URL"""
        if not HAS_DOCX:
            return {"url": url, "error": "DOCX解析功能不可用，请安装 python-docx", "error_code": ErrorCode.DOCX_PARSE_ERROR.value}
        
        try:
            response = self.session.get(url, timeout=60)
            response.raise_for_status()
            
            with tempfile.NamedTemporaryFile(delete=False, suffix='.docx') as tmp_file:
                tmp_file.write(response.content)
                tmp_path = tmp_file.name
            
            doc = docx.Document(tmp_path)
            content = "\n".join([para.text for para in doc.paragraphs])
            
            os.unlink(tmp_path)
            
            return {
                "url": url,
                "title": url.split('/')[-1],
                "content": content[:10000],
                "format": "docx"
            }
            
        except Exception as e:
            return {"url": url, "error": f"DOCX解析失败: {str(e)}"}
    
    def _parse_xlsx_url(self, url: str) -> Dict:
        """解析XLSX URL"""
        if not HAS_XLSX:
            return {"url": url, "error": "XLSX解析功能不可用，请安装 openpyxl", "error_code": ErrorCode.XLSX_PARSE_ERROR.value}
        
        try:
            response = self.session.get(url, timeout=60)
            response.raise_for_status()
            
            with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_file:
                tmp_file.write(response.content)
                tmp_path = tmp_file.name
            
            wb = openpyxl.load_workbook(tmp_path)
            content = ""
            for sheet in wb.worksheets:
                content += f"=== Sheet: {sheet.title} ===\n"
                for row in sheet.iter_rows(max_row=100, values_only=True):
                    row_text = " | ".join([str(cell) if cell else "" for cell in row])
                    if row_text.strip():
                        content += row_text + "\n"
            
            os.unlink(tmp_path)
            
            return {
                "url": url,
                "title": url.split('/')[-1],
                "content": content[:10000],
                "format": "xlsx"
            }
            
        except Exception as e:
            return {"url": url, "error": f"XLSX解析失败: {str(e)}"}
    
    def extract_urls_from_text(self, text: str) -> List[str]:
        """从文本中提取URL"""
        url_pattern = r'https?://[^\s<>"{}|\\^`\[\]]+'
        urls = re.findall(url_pattern, text)
        return list(set(urls))


web_parser = WebParser() if HAS_WEB_PARSER else None


def get_time_context():
    """获取当前时间上下文"""
    now = datetime.now()
    weekdays = ['一', '二', '三', '四', '五', '六', '日']
    
    year = now.year
    is_leap = (year % 4 == 0 and year % 100 != 0) or (year % 400 == 0)
    
    month_names = ['', '一月', '二月', '三月', '四月', '五月', '六月', 
                   '七月', '八月', '九月', '十月', '十一月', '十二月']
    days_in_month = [0, 31, 29 if is_leap else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    
    return f"""【当前时间信息】
当前日期：{now.year}年{now.month}月{now.day}日
星期：星期{weekdays[now.weekday()]}
当前时间：{now.hour}时{now.minute}分
年份：{now.year}年
是否闰年：{"是" if is_leap else "否"}（{year}年{"有366天" if is_leap else "有365天"}）
本月：{month_names[now.month]}，共{days_in_month[now.month]}天
下个月：{month_names[now.month + 1] if now.month < 12 else "一月"}，共{days_in_month[now.month + 1] if now.month < 12 else days_in_month[1]}天
"""


class KnowledgeManager:
    """知识管理器"""
    
    def __init__(self, knowledge_base_path, vector_store_path):
        self.local_rag = LocalRAG(knowledge_base_path)
        self.vector_rag = VectorRAG(vector_store_path)
        self.ai_personality_manager = None
        self._cached_url_results = []
        self._cached_url_context = ""
    
    def set_ai_personality_manager(self, manager):
        """设置AI人格管理器"""
        self.ai_personality_manager = manager
    
    def add_knowledge(self, question, answer):
        """添加知识"""
        return self.local_rag.add_knowledge(question, answer)
    
    def search_knowledge(self, query, threshold=0.5):
        """搜索知识"""
        local_results = self.local_rag.search_knowledge(query, threshold)
        query_vector = self._generate_simple_vector(query)
        vector_results = self.vector_rag.search_vectors(query_vector)
        
        results = local_results
        for item in vector_results:
            if item not in results:
                results.append(item)
        return results[:5]
    
    def _build_system_prompt(self) -> str:
        """构建系统提示词"""
        time_context = get_time_context()
        
        ai_prompt = self.ai_personality_manager.get_system_prompt() if self.ai_personality_manager else ""
        
        system_prompt = f"""{ai_prompt}

{time_context}

【回答要求】
- 保持AI人格设定的性格特点
- 回答要详细、有条理，包含实际建议和例子
- 如果用户问日期、时间、闰年、月份天数等问题，请直接使用时间信息回答
- 记住用户的重要信息，在后续对话中体现"""
        
        return system_prompt
    
    def search_with_api(self, query, memory_results=[], conversation_history=None, max_retries=3):
        """使用 API 搜索（流式输出，带重试机制）
        
        Args:
            query: 用户问题
            memory_results: 相关记忆结果
            conversation_history: 对话历史列表，每项包含 user_input 和 system_response
            max_retries: 最大重试次数
        """
        if not HAS_API:
            yield "[系统提示] API 依赖未安装，无法使用 API 搜索"
            return
        
        system_prompt = self._build_system_prompt()
        
        messages = [{"role": "system", "content": system_prompt}]
        
        if conversation_history:
            for turn in conversation_history[-5:]:
                if turn.get("user_input"):
                    messages.append({"role": "user", "content": turn["user_input"]})
                if turn.get("system_response"):
                    messages.append({"role": "assistant", "content": turn["system_response"]})
        
        user_content_parts = []
        
        if memory_results:
            memory_prompt = "【相关记忆】" + "; ".join(memory_results[:2])[:300]
            user_content_parts.append(memory_prompt)
        
        user_content_parts.append(f"问题：{query}")
        user_content_parts.append("请提供详细、有帮助的回答。如果这是追问，请结合之前的对话内容回答。")
        
        messages.append({"role": "user", "content": "\n\n".join(user_content_parts)})
        
        for attempt in range(max_retries):
            try:
                stream = client.chat.completions.create(
                    model=MODEL_ID,
                    messages=messages,
                    temperature=0.7,
                    stream=True,
                    timeout=30,
                    max_tokens=800
                )
                
                for chunk in stream:
                    if chunk.choices[0].delta.content is not None:
                        yield chunk.choices[0].delta.content
                
                return
                
            except Exception as e:
                error_msg = str(e)
                
                if "timeout" in error_msg.lower() or "timed out" in error_msg.lower():
                    if attempt < max_retries - 1:
                        yield f"\n[系统提示] 网络超时，正在重试 ({attempt + 1}/{max_retries})...\n"
                        import time
                        time.sleep(1)
                        continue
                    else:
                        yield "\n[系统提示] 抱歉，网络连接超时，请检查网络后重试。"
                        return
                else:
                    yield f"\n[系统提示] 服务暂时不可用，请稍后再试。"
                    return
    
    def search_with_knowledge_context(self, query, knowledge_context, memory_results=[], conversation_history=None, max_retries=3):
        """带知识库上下文的 API 搜索（流式输出）

        将知识库匹配结果作为参考上下文注入 LLM，让 AI 在此基础上补充完善回答。

        Args:
            query: 用户问题
            knowledge_context: 知识库匹配到的参考内容
            memory_results: 相关记忆结果
            conversation_history: 对话历史列表
            max_retries: 最大重试次数
        """
        if not HAS_API:
            yield knowledge_context or "[系统提示] API 依赖未安装，无法使用 API 搜索"
            return

        system_prompt = self._build_system_prompt()

        messages = [{"role": "system", "content": system_prompt}]

        if conversation_history:
            for turn in conversation_history[-5:]:
                if turn.get("user_input"):
                    messages.append({"role": "user", "content": turn["user_input"]})
                if turn.get("system_response"):
                    messages.append({"role": "assistant", "content": turn["system_response"]})

        user_content_parts = []

        if memory_results:
            memory_prompt = "【相关记忆】" + "; ".join(memory_results[:2])[:300]
            user_content_parts.append(memory_prompt)

        if knowledge_context:
            knowledge_prompt = f"【知识库参考信息】\n{knowledge_context[:500]}\n\n请参考以上知识库信息，结合你自身的知识，给出更完整、详细、有帮助的回答。如果知识库信息不够充分，请主动补充相关内容。"
            user_content_parts.append(knowledge_prompt)

        user_content_parts.append(f"问题：{query}")

        messages.append({"role": "user", "content": "\n\n".join(user_content_parts)})

        for attempt in range(max_retries):
            try:
                stream = client.chat.completions.create(
                    model=MODEL_ID,
                    messages=messages,
                    temperature=0.7,
                    stream=True,
                    timeout=30,
                    max_tokens=1000
                )

                for chunk in stream:
                    if chunk.choices[0].delta.content is not None:
                        yield chunk.choices[0].delta.content

                return

            except Exception as e:
                error_msg = str(e)

                if "timeout" in error_msg.lower() or "timed out" in error_msg.lower():
                    if attempt < max_retries - 1:
                        yield f"\n[系统提示] 网络超时，正在重试 ({attempt + 1}/{max_retries})...\n"
                        import time
                        time.sleep(1)
                        continue
                    else:
                        yield "\n[系统提示] 抱歉，网络连接超时，请检查网络后重试。"
                        return
                else:
                    yield f"\n[系统提示] 服务暂时不可用，请稍后再试。"
                    return

    def _generate_simple_vector(self, text):
        """生成简单的向量表示（仅用于演示）"""
        vector = []
        for char in 'abcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ':
            vector.append(text.count(char) / len(text) if text else 0)
        return vector
    
    def get_knowledge_count(self):
        """获取知识数量"""
        return self.local_rag.get_knowledge_count()
    
    def clear_knowledge_base(self):
        """清空知识库"""
        self.local_rag.clear_knowledge_base()
        self.vector_rag.clear_vector_store()
    
    def add_vector(self, text, vector):
        """添加向量"""
        return self.vector_rag.add_vector(text, vector)
    
    def parse_urls(self, url_list: List[str]) -> List[Dict]:
        """解析URL列表（最多3个）"""
        if not web_parser:
            return [{"error": "网页解析功能不可用，请安装 requests 和 beautifulsoup4"}]
        return web_parser.parse_urls(url_list)
    
    def parse_single_url(self, url: str) -> Dict:
        """解析单个URL"""
        if not web_parser:
            return {"error": "网页解析功能不可用"}
        return web_parser.parse_single_url(url)
    
    def extract_urls_from_text(self, text: str) -> List[str]:
        """从文本中提取URL"""
        if not web_parser:
            return []
        return web_parser.extract_urls_from_text(text)
    
    def parse_urls_from_query(self, query: str) -> Tuple[str, List[Dict]]:
        """从查询中提取并解析URL，返回清理后的查询和解析结果"""
        urls = self.extract_urls_from_text(query)
        if not urls:
            return query, []
        
        results = self.parse_urls(urls)
        
        cleaned_query = query
        for url in urls:
            cleaned_query = cleaned_query.replace(url, '').strip()
        
        return cleaned_query, results
    
    def search_with_url_context(self, query: str, memory_results: List[str] = [], conversation_history=None, max_retries: int = 3):
        """带URL上下文的API搜索（流式输出）
        
        Args:
            query: 用户问题
            memory_results: 相关记忆结果
            conversation_history: 对话历史列表
            max_retries: 最大重试次数
        """
        if not HAS_API:
            yield "[系统提示] API 依赖未安装，无法使用 API 搜索"
            return
        
        cleaned_query, url_results = self.parse_urls_from_query(query)
        
        if url_results:
            self._cached_url_results = url_results
            self._cached_url_context = ""
            for result in url_results:
                if "error" in result:
                    self._cached_url_context += f"URL: {result.get('url', '未知')} - 解析失败: {result['error']}\n"
                else:
                    self._cached_url_context += f"URL: {result.get('url', '')}\n"
                    self._cached_url_context += f"标题: {result.get('title', '无标题')}\n"
                    self._cached_url_context += f"内容: {result.get('content', '')}\n\n"
        
        system_prompt = self._build_system_prompt()
        
        messages = [{"role": "system", "content": system_prompt}]
        
        if conversation_history:
            for turn in conversation_history[-5:]:
                if turn.get("user_input"):
                    messages.append({"role": "user", "content": turn["user_input"]})
                if turn.get("system_response"):
                    messages.append({"role": "assistant", "content": turn["system_response"]})
        
        user_content_parts = []
        
        url_context = f"【网页解析内容】\n{self._cached_url_context}" if self._cached_url_context else ""
        
        if memory_results:
            memory_prompt = "【相关记忆】" + "; ".join(memory_results[:2])[:300]
            user_content_parts.append(memory_prompt)
        
        if url_context:
            user_content_parts.append(url_context)
        
        if cleaned_query:
            user_content_parts.append(f"问题：{cleaned_query}")
        elif self._cached_url_context:
            user_content_parts.append("请根据以上网页内容进行总结和分析。")
        else:
            user_content_parts.append(f"问题：{query}")
        
        user_content_parts.append("请提供详细、有帮助的回答。如果这是追问，请结合之前的对话内容回答。")
        
        messages.append({"role": "user", "content": "\n\n".join(user_content_parts)})
        
        for attempt in range(max_retries):
            try:
                stream = client.chat.completions.create(
                    model=MODEL_ID,
                    messages=messages,
                    temperature=0.7,
                    stream=True,
                    timeout=60,
                    max_tokens=1500
                )
                
                for chunk in stream:
                    if chunk.choices[0].delta.content is not None:
                        yield chunk.choices[0].delta.content
                
                return
                
            except Exception as e:
                error_msg = str(e)
                
                if "timeout" in error_msg.lower() or "timed out" in error_msg.lower():
                    if attempt < max_retries - 1:
                        yield f"\n[系统提示] 网络超时，正在重试 ({attempt + 1}/{max_retries})...\n"
                        import time
                        time.sleep(1)
                        continue
                    else:
                        yield "\n[系统提示] 抱歉，网络连接超时，请检查网络后重试。"
                        return
                else:
                    yield f"\n[系统提示] 服务暂时不可用，请稍后再试。"
                    return
    
    def clear_url_cache(self):
        """清除网页内容缓存"""
        self._cached_url_results = []
        self._cached_url_context = ""
    
    def has_cached_url(self) -> bool:
        """检查是否有缓存的网页内容"""
        return bool(self._cached_url_context)
